"""Lager excel/LMS-timer.xlsx: kalender, timeføring og rapport som en Excel-arbeidsbok.

Samme regler som appen (src/model.js, src/schedule.js, src/state.js):
skiftplan på 21 dager, arbeidstider, faste aktiviteter/maskiner/arbeid og LMS-standard
(arbeidsdager uten nok føringer fylles automatisk med LMS, rundet til nærmeste halvtime).

Bare anonyme eksempelnavn («Tekniker 1-01» osv.). Kjør:  python3 excel/lag_excel.py
"""
import datetime as dt
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import DataBarRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as col
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

OUT = Path(__file__).with_name('LMS-timer.xlsx')

# ---------- domene (speiler src/model.js) ----------
SHIFTS = [(1, 'Skift 1', dt.date(2026, 1, 5)), (2, 'Skift 2', dt.date(2026, 1, 12)), (3, 'Skift 3', dt.date(2026, 1, 19))]
TEAM = {1: 11, 2: 12, 3: 11}
SLOTS = 18                      # plasser per skift i Oppsett
NSLOT = SLOTS * len(SHIFTS)
# (navn, kode i kalenderen, rolle, farge --aN)
ACTS = [('Line', 'LN', '', '2A6FDB'), ('LMS', 'LMS', '', '5C7383'), ('Assessor', 'A', 'Assessor', '7457C9'),
        ('Assessment', 'AS', '', '8A5A44'), ('Hot Section', 'HS', 'Hot Section', 'C9372C'), ('Kurs', 'KU', '', '14929A'),
        ('Mod', 'M', '', 'E0702F'), ('Ferie', 'FE', '', '2E9457'), ('Avspasering', 'AV', '', '79BD8F'),
        ('Fravær', 'FR', '', 'A0527A'), ('Annet', 'AN', '', 'B08A1E')]
MACHINES = ['OIA', 'OIB', 'OIC', 'OIE', 'OIF', 'OIG', 'OIH', 'OII', 'OIK', 'OIM', 'OIN', 'OMI', 'ONA', 'OND', 'ONE',
            'ONH', 'ONM', 'ONO', 'ONP', 'ONQ', 'ONT', 'ONW', 'ONY', 'ONS']
WORK = ['Sheet', 'Repair', 'Inspection', 'Mod', 'SLL', 'Testing', 'Final']
MONTHS = ['januar', 'februar', 'mars', 'april', 'mai', 'juni', 'juli', 'august', 'september', 'oktober', 'november', 'desember']
# fase, kode, dager, man–tor fra/til, fredag fra/til
PHASES = [('Kveld', 'K', 5, '14:30', '23:30', '10:30', '19:00'), ('Helg', 'H', 2, '07:00', '16:30', '07:00', '16:30'),
          ('Dag', 'D', 5, '06:00', '15:00', '06:00', '13:00'), ('Fri', 'F', 9, None, None, None, None)]
AUTO_FROM = dt.date(2026, 1, 1)
GRID_DAYS = 3 * 366             # Grunnlag dekker tre år fra AutoFra
ROWS = 3000                     # føringsrader i Timeføring

# ---------- stil (designtokens fra styles.css) ----------
FONT = 'Arial'
F = lambda **k: Font(name=FONT, size=k.pop('size', 10), **k)
fill = lambda c: PatternFill('solid', start_color=c, end_color=c)
LINE = Side(style='thin', color='E1E4E8')
BOX = Border(left=LINE, right=LINE, top=LINE, bottom=LINE)
HEAD = fill('F7F8FA'); INPUT = fill('E8F0FB'); ACCENT = '1F5FBF'
PH_FILL = {'K': 'DDD5F7', 'H': 'FAE0B8', 'D': 'CFE4F7'}
DATE = 'dd.mm.yyyy'; HRS = '0.0;-0.0;"–"'
CENTER = Alignment(horizontal='center', vertical='center')


def head(ws, row, c1, labels, wrap=False):
    for i, t in enumerate(labels):
        c = ws.cell(row, c1 + i, t); c.font = F(bold=True); c.fill = HEAD; c.border = BOX
        c.alignment = Alignment(vertical='center', wrap_text=wrap, horizontal='left' if i == 0 else 'center')


def title(ws, text, sub):
    ws['A1'] = text; ws['A1'].font = F(size=16, bold=True)
    ws['A2'] = sub; ws['A2'].font = F(color='636B74')


def name(wb, n, ref):
    wb.defined_names[n] = DefinedName(n, attr_text=ref)


wb = Workbook()
kal = wb.active; kal.title = 'Kalender'
tf = wb.create_sheet('Timeføring'); rap = wb.create_sheet('Rapport')
opp = wb.create_sheet('Oppsett'); gr = wb.create_sheet('Grunnlag')

# =====================================================================
# Oppsett
# =====================================================================
title(opp, 'Oppsett', 'Blå celler fylles inn. Teknikere står under skiftet sitt (flytt ved å klippe/lime raden til et annet skift; føringene beholdes).')
head(opp, 4, 1, ['Skift', 'Tekniker', 'Ekstra fra', 'Ekstra til', 'Roller'])
O0 = 5                                   # første teknikerrad
for s, sname, _ in SHIFTS:
    for n in range(1, SLOTS + 1):
        r = O0 + (s - 1) * SLOTS + n - 1
        opp.cell(r, 1, f'=INDEX(SkiftNavn,{s})').font = F(color='636B74')
        for c in range(2, 6):
            cell = opp.cell(r, c); cell.fill = INPUT; cell.border = BOX; cell.font = F()
            if c in (3, 4): cell.number_format = DATE
        if n <= TEAM[s]:
            opp.cell(r, 2, f'Tekniker {s}-{n:02d}')
O1 = O0 + NSLOT - 1
opp.cell(O0, 5).comment = Comment('Roller skilt med komma, f.eks. «B1, Assessor». Assessor og Hot Section krever rollen.', 'LMS Hours')
opp.cell(O0, 3).comment = Comment('Ekstra personell: fyll inn fra/til. Teknikeren er da bare aktiv i perioden.', 'LMS Hours')
for c, w in zip('ABCDE', [10, 20, 12, 12, 24]): opp.column_dimensions[c].width = w
name(wb, 'Teknikere', f'Oppsett!$B${O0}:$B${O1}')
name(wb, 'TekFra', f'Oppsett!$C${O0}:$C${O1}')
name(wb, 'TekTil', f'Oppsett!$D${O0}:$D${O1}')
name(wb, 'TekRoller', f'Oppsett!$E${O0}:$E${O1}')

# skift
head(opp, 4, 7, ['Skift', 'Navn', 'Startdato'])
for i, (s, sname, start) in enumerate(SHIFTS):
    r = 5 + i
    opp.cell(r, 7, s).font = F(); opp.cell(r, 8, sname).font = F()
    c = opp.cell(r, 9, start); c.number_format = DATE; c.font = F(); c.fill = INPUT
name(wb, 'SkiftNavn', 'Oppsett!$H$5:$H$7'); name(wb, 'SkiftStart', 'Oppsett!$I$5:$I$7')

# syklus og arbeidstid
head(opp, 10, 7, ['Fase', 'Kode', 'Dager', 'Man–tor fra', 'Man–tor til', 'Fredag fra', 'Fredag til', 'Timer man–tor', 'Timer fredag'], wrap=True)
for i, (pn, code, days, a, b, c_, d) in enumerate(PHASES):
    r = 11 + i
    opp.cell(r, 7, pn); opp.cell(r, 8, code); opp.cell(r, 9, days).fill = INPUT
    for j, v in enumerate([a, b, c_, d]):
        if v:
            h, m = map(int, v.split(':')); cell = opp.cell(r, 10 + j, dt.time(h, m)); cell.number_format = 'hh:mm'; cell.fill = INPUT
    if code != 'F':
        opp.cell(r, 14, f'=MOD(K{r}-J{r},1)*24'); opp.cell(r, 15, f'=MOD(M{r}-L{r},1)*24')
    else:
        opp.cell(r, 14, 9).fill = INPUT; opp.cell(r, 15, '=N14')
        opp.cell(r, 14).comment = Comment('Overtid på fridag: 9 t (føres som vanlig, regnes som overtid).', 'LMS Hours')
    for cc in range(7, 16):
        opp.cell(r, cc).border = BOX; opp.cell(r, cc).font = F()
        if cc >= 14: opp.cell(r, cc).number_format = '0.0'
opp['G15'] = 'Syklus (dager)'; opp['I15'] = '=SUM(I11:I14)'
opp['G17'] = 'LMS-standard fra'; opp['I17'] = AUTO_FROM; opp['I17'].number_format = DATE; opp['I17'].fill = INPUT
for a in ('G15', 'I15', 'G17', 'I17'): opp[a].font = F(bold=a.startswith('G'))
name(wb, 'FaseKode', 'Oppsett!$H$11:$H$13'); name(wb, 'TimerMT', 'Oppsett!$N$11:$N$13'); name(wb, 'TimerFr', 'Oppsett!$O$11:$O$13')
for n_, r in (('DagK', 11), ('DagH', 12), ('DagD', 13)): name(wb, n_, f'Oppsett!$I${r}')
name(wb, 'Syklus', 'Oppsett!$I$15'); name(wb, 'AutoFra', 'Oppsett!$I$17')
opp['G19'] = 'LMS-standard: arbeidsdager fylles automatisk med LMS for det som er igjen av dagen etter føringene i Timeføring (rundet til nærmeste halvtime), til og med i dag.'
opp['G19'].font = F(color='636B74', size=9)

# faste lister
head(opp, 4, 17, ['Aktivitet', 'Kode', 'Krever rolle'])
for i, (an, code, role, _) in enumerate(ACTS):
    for j, v in enumerate([an, code, role]):
        c = opp.cell(5 + i, 17 + j, v); c.font = F(); c.border = BOX
name(wb, 'Aktiviteter', f'Oppsett!$Q$5:$Q${4 + len(ACTS)}'); name(wb, 'AktKode', f'Oppsett!$R$5:$R${4 + len(ACTS)}')
name(wb, 'AktRolle', f'Oppsett!$S$5:$S${4 + len(ACTS)}')
opp.cell(5 + len(ACTS), 17, 'Faste aktiviteter (endres ikke).').font = F(color='636B74', size=9)

head(opp, 4, 21, ['Maskin', 'Status'])
for i, m in enumerate(MACHINES):
    opp.cell(5 + i, 21, m).font = F(); c = opp.cell(5 + i, 22, 'Ferdig'); c.fill = INPUT; c.font = F()
    opp.cell(5 + i, 21).border = BOX; c.border = BOX
M1 = 4 + len(MACHINES)
name(wb, 'Maskiner', f'Oppsett!$U$5:$U${M1}'); name(wb, 'MaskinStatus', f'Oppsett!$V$5:$V${M1}')
dv = DataValidation(type='list', formula1='"I dokk,Ferdig"', allow_blank=True); opp.add_data_validation(dv); dv.add(f'V5:V{M1}')
opp.cell(M1 + 1, 21, 'LMS og Mod: bare maskiner «I dokk». Line: alle.').font = F(color='636B74', size=9)

head(opp, 4, 24, ['Arbeid (LMS)'])
for i, w in enumerate(WORK): opp.cell(5 + i, 24, w).font = F(); opp.cell(5 + i, 24).border = BOX
name(wb, 'Arbeid', f'Oppsett!$X$5:$X${4 + len(WORK)}')
head(opp, 4, 26, ['Måned'])
for i, m in enumerate(MONTHS): opp.cell(5 + i, 26, m).font = F()
name(wb, 'Maaneder', 'Oppsett!$Z$5:$Z$16')
for c, w in {'F': 3, 'G': 15, 'H': 9, 'I': 11, 'J': 9, 'K': 9, 'L': 9, 'M': 9, 'N': 9, 'O': 9, 'P': 3, 'Q': 14, 'R': 7,
             'S': 13, 'T': 3, 'U': 9, 'V': 9, 'W': 3, 'X': 13, 'Y': 3, 'Z': 11}.items():
    opp.column_dimensions[c].width = w
opp.row_dimensions[10].height = 28
opp.freeze_panes = 'A5'

# =====================================================================
# Timeføring
# =====================================================================
title(tf, 'Timeføring', 'Én rad per føring. Fyll inn A–G (blå overskrift); resten regnes ut. Dager uten nok føringer fylles med LMS automatisk.')
T0, T1 = 4, 3 + ROWS
cols = ['Dato', 'Tekniker', 'Aktivitet', 'Maskin', 'Arbeid', 'Timer', 'Notat', 'Skift', 'Fase', 'Overtid', 'Kontroll',
        'SkiftNr', 'Nøkkel', 'AktSum', 'Dominant']
head(tf, 3, 1, cols)
for c in range(1, 8): tf.cell(3, c).fill = INPUT
for c, w in zip('ABCDEFGHIJK', [11, 18, 13, 9, 11, 7, 28, 9, 8, 8, 26]): tf.column_dimensions[c].width = w
for c in 'LMNO': tf.column_dimensions[c].hidden = True
rng = lambda c: f'Timeføring!${c}${T0}:${c}${T1}'
for n_, c in (('TF_Dato', 'A'), ('TF_Tekniker', 'B'), ('TF_Akt', 'C'), ('TF_Maskin', 'D'), ('TF_Arbeid', 'E'), ('TF_Timer', 'F'),
              ('TF_Overtid', 'J'), ('TF_Nokkel', 'M'), ('TF_AktSum', 'N'), ('TF_Dom', 'O')):
    name(wb, n_, rng(c))
phase_name = lambda d, s: (f'IF(MOD({d}-INDEX(SkiftStart,{s}),Syklus)<DagK,"Kveld",IF(MOD({d}-INDEX(SkiftStart,{s}),Syklus)<DagK+DagH,"Helg",'
                           f'IF(MOD({d}-INDEX(SkiftStart,{s}),Syklus)<DagK+DagH+DagD,"Dag","Fri")))')
for r in range(T0, T1 + 1):
    tf.cell(r, 1).number_format = DATE; tf.cell(r, 6).number_format = '0.0'
    tf[f'L{r}'] = f'=IF(B{r}="","",IFERROR(INT((MATCH(B{r},Teknikere,0)-1)/{SLOTS})+1,""))'
    tf[f'H{r}'] = f'=IF(L{r}="","",INDEX(SkiftNavn,L{r}))'
    tf[f'I{r}'] = f'=IF(OR(A{r}="",L{r}=""),"",{phase_name(f"A{r}", f"L{r}")})'
    tf[f'J{r}'] = f'=IF(I{r}="Fri","Ja","")'
    rolle = f'INDEX(AktRolle,MATCH(C{r},Aktiviteter,0))&""'
    tf[f'K{r}'] = (f'=IF(B{r}="","",IF(L{r}="","Ukjent tekniker",IF(ISNA(MATCH(C{r},Aktiviteter,0)),"Ukjent aktivitet",'
                   f'IF(AND({rolle}<>"",ISERROR(SEARCH({rolle},INDEX(TekRoller,MATCH(B{r},Teknikere,0))))),"Mangler rollen "&{rolle},'
                   f'IF(AND(C{r}="LMS",F{r}>0,OR(D{r}="",E{r}="")),"Mangler maskin eller arbeid",'
                   f'IF(AND(OR(C{r}="LMS",C{r}="Mod"),D{r}<>""),IF(IFERROR(INDEX(MaskinStatus,MATCH(D{r},Maskiner,0)),"")="I dokk","","Maskinen er ikke i dokk"),""))))))')
    tf[f'M{r}'] = f'=IF(OR(A{r}="",B{r}=""),"",B{r}&"|"&A{r})'
    tf[f'N{r}'] = f'=IF(M{r}="","",SUMIFS(TF_Timer,TF_Nokkel,M{r},TF_Akt,C{r}))'
    tf[f'O{r}'] = (f'=IF(M{r}="","",IF(AND(N{r}=_xlfn.MAXIFS(TF_AktSum,TF_Nokkel,M{r}),'
                   f'COUNTIFS(M${T0 - 1}:M{r - 1},M{r},N${T0 - 1}:N{r - 1},N{r})=0),M{r},""))')
    for c in range(1, 12): tf.cell(r, c).font = F(color='636B74' if c >= 8 else '1B1F24')
# eksempelrad
for c, v in enumerate([dt.date(2026, 1, 6), 'Tekniker 1-01', 'Kurs', None, None, 4, 'Eksempel – slett denne raden'], 1):
    tf.cell(T0, c, v)
tf.cell(T0, 1).number_format = DATE
for c, src in ((2, 'Teknikere'), (3, 'Aktiviteter'), (4, 'Maskiner'), (5, 'Arbeid')):
    dv = DataValidation(type='list', formula1=f'={src}', allow_blank=True); tf.add_data_validation(dv); dv.add(f'{col(c)}{T0}:{col(c)}{T1}')
dv = DataValidation(type='date', operator='greaterThan', formula1='36526', allow_blank=True, error='Skriv en dato, f.eks. 06.01.2026')
tf.add_data_validation(dv); dv.add(f'A{T0}:A{T1}')
dv = DataValidation(type='decimal', operator='between', formula1='0', formula2='24', allow_blank=True, error='Timer mellom 0 og 24')
tf.add_data_validation(dv); dv.add(f'F{T0}:F{T1}')
tf.conditional_formatting.add(f'K{T0}:K{T1}', FormulaRule(formula=[f'K{T0}<>""'], font=Font(color='B25E00', bold=True), fill=fill('FFF4E5')))
tf.conditional_formatting.add(f'J{T0}:J{T1}', FormulaRule(formula=[f'J{T0}="Ja"'], font=Font(color='B25E00', bold=True)))
tf.auto_filter.ref = f'A3:K{T1}'
tf.freeze_panes = 'A4'

# =====================================================================
# Grunnlag (skjult): fase, normal arbeidsdag og automatisk LMS per tekniker og dag
# =====================================================================
G0 = 9; G1 = G0 + GRID_DAYS - 1; GC = 8          # første datorad, siste datorad, første teknikerkolonne
gr['A1'] = 'Grunnlag (hjelpeark – ikke endre)'; gr['A1'].font = F(bold=True)
gr['A2'] = 'Automatisk LMS (LMS-standard) per tekniker og dag. Brukes av Kalender og Rapport.'; gr['A2'].font = F(color='636B74')
for lab, r in (('Plass', 3), ('Tekniker', 4), ('Skift', 5), ('Fra', 6), ('Til', 7)): gr.cell(r, 7, lab).font = F(bold=True)
head(gr, 8, 1, ['Dato', 'Fase S1', 'Fase S2', 'Fase S3', 'Timer S1', 'Timer S2', 'Timer S3'])
for k in range(1, NSLOT + 1):
    c = col(GC + k - 1)
    gr[f'{c}3'] = k; gr[f'{c}4'] = f'=INDEX(Teknikere,{c}3)&""'; gr[f'{c}5'] = (k - 1) // SLOTS + 1
    gr[f'{c}6'] = f'=INDEX(TekFra,{c}3)'; gr[f'{c}7'] = f'=INDEX(TekTil,{c}3)'
phase_key = lambda d, s: (f'IF(MOD({d}-INDEX(SkiftStart,{s}),Syklus)<DagK,"K",IF(MOD({d}-INDEX(SkiftStart,{s}),Syklus)<DagK+DagH,"H",'
                          f'IF(MOD({d}-INDEX(SkiftStart,{s}),Syklus)<DagK+DagH+DagD,"D","F")))')
std = lambda d, p: f'IF(WEEKDAY({d},2)=5,INDEX(TimerFr,MATCH({p},FaseKode,0)),INDEX(TimerMT,MATCH({p},FaseKode,0)))'
LAST = col(GC + NSLOT - 1)
for r in range(G0, G1 + 1):
    gr[f'A{r}'] = '=AutoFra' if r == G0 else f'=A{r - 1}+1'; gr[f'A{r}'].number_format = DATE
    for s in (1, 2, 3):
        p = col(1 + s); gr[f'{p}{r}'] = f'={phase_key(f"$A{r}", s)}'
        gr[f'{col(4 + s)}{r}'] = f'=IF({p}{r}="F",0,{std(f"$A{r}", f"{p}{r}")})'
    for k in range(1, NSLOT + 1):
        c = col(GC + k - 1)
        gr[f'{c}{r}'] = (f'=IF(OR({c}$4="",$A{r}>TODAY()),0,IF(OR(INDEX($B{r}:$D{r},{c}$5)="F",AND({c}$6>0,$A{r}<{c}$6),AND({c}$7>0,$A{r}>{c}$7)),0,'
                         f'MAX(0,ROUND((INDEX($E{r}:$G{r},{c}$5)-SUMIFS(TF_Timer,TF_Tekniker,{c}$4,TF_Dato,$A{r}))*2,0)/2)))')
name(wb, 'GR_Dato', f'Grunnlag!$A${G0}:$A${G1}'); name(wb, 'GR_Auto', f'Grunnlag!${col(GC)}${G0}:${LAST}${G1}')
gr.freeze_panes = f'{col(GC)}{G0}'
gr.sheet_state = 'hidden'

# =====================================================================
# Kalender
# =====================================================================
DC = 3                                            # første dagkolonne (C)
DAYCOLS = [col(DC + i) for i in range(31)]
DL = DAYCOLS[-1]
kal['A1'] = 'Kalender'; kal['A1'].font = F(size=16, bold=True)
kal['A2'] = 'År'; kal['B2'] = '=YEAR(TODAY())'; kal['D2'] = 'Måned'; kal['H2'] = '=INDEX(Maaneder,MONTH(TODAY()))'
for a in ('B2', 'H2'): kal[a].fill = INPUT; kal[a].font = F(bold=True); kal[a].border = BOX
kal.merge_cells('D2:G2'); kal.merge_cells('H2:L2')
for a in ('A2', 'D2'): kal[a].font = F(color='636B74'); kal[a].alignment = Alignment(horizontal='right')
dv = DataValidation(type='list', formula1='=Maaneder'); kal.add_data_validation(dv); dv.add('H2')
kal['N2'] = 'Velg år og måned. Føringer gjøres i Timeføring. «+» = flere aktiviteter samme dag. Lys LMS = forhåndsvisning (fremtidige dager).'
kal['N2'].font = F(color='636B74', size=9)
# skjulte hjelperader 7–13: første dag, fase og normaltimer per skift
kal['A7'] = f'=DATE($B$2,MATCH($H$2,Maaneder,0),1)'
for i, c in enumerate(DAYCOLS):
    kal[f'{c}5'] = f'=IF({i + 1}<=DAY(EOMONTH($A$7,0)),$A$7+{i},"")'
    kal[f'{c}4'] = f'=IF({c}5="","",IF(OR(WEEKDAY({c}5,2)=1,{i}=0),"Uke "&_xlfn.ISOWEEKNUM({c}5),""))'
    kal[f'{c}6'] = f'=IF({c}5="","",CHOOSE(WEEKDAY({c}5,2),"ma","ti","on","to","fr","lø","sø"))'
    for s in (1, 2, 3):
        kal[f'{c}{7 + s}'] = f'=IF({c}$5="","",IF({phase_key(f"{c}$5", s)}="F","",{phase_key(f"{c}$5", s)}))'
        kal[f'{c}{10 + s}'] = f'=IF({c}{7 + s}="","",{std(f"{c}$5", f"{c}{7 + s}")})'
    for r in (4, 5, 6):
        kal[f'{c}{r}'].alignment = CENTER; kal[f'{c}{r}'].font = F(size=8, color='636B74', bold=(r == 5))
    kal[f'{c}5'].number_format = 'd'; kal[f'{c}5'].font = F(size=10, bold=True)
for r in range(7, 14): kal.row_dimensions[r].hidden = True
kal['A6'] = 'Tekniker'; kal['A6'].font = F(bold=True)

KR = {}                                           # første teknikerrad per skift
r = 15
for s, sname, _ in SHIFTS:
    kal[f'A{r}'] = f'=INDEX(SkiftNavn,{s})'; kal[f'A{r}'].font = F(bold=True)
    for c in DAYCOLS:
        kal[f'{c}{r}'] = f'={c}{7 + s}'; kal[f'{c}{r}'].alignment = CENTER; kal[f'{c}{r}'].font = F(size=8, bold=True, color='2B2F36')
    kal.conditional_formatting.add(f'{DAYCOLS[0]}{r}:{DL}{r}', FormulaRule(formula=[f'{DAYCOLS[0]}{r}="K"'], fill=fill(PH_FILL['K'])))
    kal.conditional_formatting.add(f'{DAYCOLS[0]}{r}:{DL}{r}', FormulaRule(formula=[f'{DAYCOLS[0]}{r}="H"'], fill=fill(PH_FILL['H'])))
    kal.conditional_formatting.add(f'{DAYCOLS[0]}{r}:{DL}{r}', FormulaRule(formula=[f'{DAYCOLS[0]}{r}="D"'], fill=fill(PH_FILL['D'])))
    r += 1; KR[s] = r
    for n in range(1, SLOTS + 1):
        k = (s - 1) * SLOTS + n
        kal[f'A{r}'] = f'=INDEX(Teknikere,{k})&""'; kal[f'A{r}'].font = F()
        kal[f'B{r}'] = f'=IF(INDEX(TekFra,{k})>0,"ekstra","")'; kal[f'B{r}'].font = F(size=8, color='B25E00')
        for c in DAYCOLS:
            d = f'{c}$5'; key = f'$A{r}&"|"&{d}'; ph = f'{c}${7 + s}'
            man = f'SUMIFS(TF_Timer,TF_Nokkel,{key})'
            auto = f'IF({ph}="",0,MAX(0,ROUND(({c}${10 + s}-{man})*2,0)/2))'
            m = f'MATCH({key},TF_Dom,0)'
            dom, domakt = f'INDEX(TF_AktSum,{m})', f'INDEX(TF_Akt,{m})'
            lms = f'SUMIFS(TF_Timer,TF_Nokkel,{key},TF_Akt,"LMS")+{auto}'
            code = f'IF({lms}>{dom},"LMS",INDEX(AktKode,MATCH({domakt},Aktiviteter,0)))'
            split = f'OR(COUNTIFS(TF_Nokkel,{key},TF_Akt,"<>"&{domakt})>0,AND({auto}>0,{domakt}<>"LMS"))'
            kal[f'{c}{r}'] = (f'=IF(OR($A{r}="",{d}=""),"",IF(OR(AND(INDEX(TekFra,{k})>0,{d}<INDEX(TekFra,{k})),AND(INDEX(TekTil,{k})>0,{d}>INDEX(TekTil,{k}))),"",'
                              f'IF(COUNTIFS(TF_Nokkel,{key})=0,IF({ph}<>"","LMS",""),IFERROR({code}&IF({split},"+",""),"?"))))')
            kal[f'{c}{r}'].alignment = CENTER; kal[f'{c}{r}'].font = F(size=8, bold=True)
        r += 1
    r += 1
KEND = r - 2
grid = f'{DAYCOLS[0]}{KR[1]}:{DL}{KEND}'
g0 = f'{DAYCOLS[0]}{KR[1]}'
kal.conditional_formatting.add(grid, FormulaRule(formula=[f'AND({DAYCOLS[0]}$5>TODAY(),{g0}="LMS")'], fill=fill('C2CCD3'), font=Font(color='2B2F36'), stopIfTrue=True))
for an, code, _, color in ACTS:
    fg = '1B1F24' if color == '79BD8F' else 'FFFFFF'
    kal.conditional_formatting.add(grid, FormulaRule(formula=[f'OR({g0}="{code}",{g0}="{code}+")'], fill=fill(color), font=Font(color=fg), stopIfTrue=True))
kal.conditional_formatting.add(grid, FormulaRule(formula=[f'AND($A{KR[1]}<>"",{DAYCOLS[0]}$5<>"")'], border=BOX))   # ruter bare på brukte rader
hdr = f'{DAYCOLS[0]}5:{DL}6'
kal.conditional_formatting.add(hdr, FormulaRule(formula=[f'{DAYCOLS[0]}$5=TODAY()'], fill=fill(ACCENT), font=Font(color='FFFFFF'), stopIfTrue=True))
kal.conditional_formatting.add(hdr, FormulaRule(formula=[f'AND({DAYCOLS[0]}$5<>"",WEEKDAY({DAYCOLS[0]}$5,2)>5)'], fill=fill('E1E4E8')))
# forklaring
LX = col(DC + 32); LY = col(DC + 33); L0 = KR[1]      # under de skjulte hjelperadene
kal[f'{LX}{L0 - 1}'] = 'Forklaring'; kal[f'{LX}{L0 - 1}'].font = F(bold=True)
for i, (an, code, _, color) in enumerate(ACTS):
    c = kal[f'{LX}{L0 + i}']; c.value = code; c.fill = fill(color); c.alignment = CENTER
    c.font = F(size=8, bold=True, color='1B1F24' if color == '79BD8F' else 'FFFFFF')
    kal[f'{LY}{L0 + i}'] = an; kal[f'{LY}{L0 + i}'].font = F(size=9)
i = L0 + len(ACTS)
kal[f'{LX}{i}'] = 'LMS'; kal[f'{LX}{i}'].fill = fill('C2CCD3'); kal[f'{LX}{i}'].font = F(size=8, bold=True); kal[f'{LX}{i}'].alignment = CENTER
kal[f'{LY}{i}'] = 'LMS (forhåndsvisning)'; kal[f'{LY}{i}'].font = F(size=9)
for j, (pn, code) in enumerate((('Kveld', 'K'), ('Helg', 'H'), ('Dag', 'D'))):
    c = kal[f'{LX}{i + 2 + j}']; c.value = code; c.fill = fill(PH_FILL[code]); c.alignment = CENTER; c.font = F(size=8, bold=True)
    kal[f'{LY}{i + 2 + j}'] = pn; kal[f'{LY}{i + 2 + j}'].font = F(size=9)
kal.column_dimensions['A'].width = 16; kal.column_dimensions['B'].width = 6
for c in DAYCOLS: kal.column_dimensions[c].width = 4.6
kal.column_dimensions[col(DC + 31)].width = 2; kal.column_dimensions[LX].width = 5; kal.column_dimensions[LY].width = 20
kal.freeze_panes = f'{DAYCOLS[0]}7'
kal.page_setup.orientation = 'landscape'; kal.page_setup.fitToWidth = 1; kal.page_setup.fitToHeight = 0
kal.sheet_properties.pageSetUpPr.fitToPage = True

# =====================================================================
# Rapport
# =====================================================================
title(rap, 'Rapport', 'Timer i perioden per aktivitet, maskin, skift og tekniker. Automatisk LMS (LMS-standard) er tatt med.')
rap['A4'] = 'Fra'; rap['B4'] = '=DATE(YEAR(TODAY()),MONTH(TODAY()),1)'
rap['A5'] = 'Til'; rap['B5'] = '=TODAY()'
for a in ('B4', 'B5'): rap[a].fill = INPUT; rap[a].number_format = DATE; rap[a].font = F(bold=True); rap[a].border = BOX
for a in ('A4', 'A5'): rap[a].font = F(bold=True)
rap['B4'].comment = Comment('Skriv inn egne datoer, eller kopier fra snarveiene til høyre.', 'LMS Hours')
name(wb, 'Fra', 'Rapport!$B$4'); name(wb, 'Til', 'Rapport!$B$5')
rap['D3'] = 'Snarveier'; rap['D3'].font = F(bold=True, color='636B74')
s1 = 'INDEX(SkiftStart,1)'
for i, (lab, f_, t_) in enumerate((('Denne måneden', 'DATE(YEAR(TODAY()),MONTH(TODAY()),1)', 'TODAY()'),
                                   ('Denne skiftperioden', f'TODAY()-MOD(TODAY()-{s1},Syklus)', f'TODAY()-MOD(TODAY()-{s1},Syklus)+Syklus-1'),
                                   ('I år', 'DATE(YEAR(TODAY()),1,1)', 'TODAY()'))):
    rap.cell(4 + i, 4, lab).font = F(color='636B74')
    for j, v in enumerate((f_, t_)):
        c = rap.cell(4 + i, 6 + j, '=' + v); c.number_format = DATE; c.font = F(color='636B74')
AUTOSUM = 'SUMPRODUCT((GR_Dato>=Fra)*(GR_Dato<=Til)*GR_Auto)'
per = 'TF_Dato,">="&Fra,TF_Dato,"<="&Til'

# nøkkeltall
A0 = 12; A1 = A0 + len(ACTS) - 1
tiles = [('A', 'Timer totalt', f'=B{A1 + 1}', HRS), ('C', 'Timer LMS', f'=B{A0 + 1}', HRS), ('E', 'LMS-andel', f'=IF(A9=0,0,C9/A9)', '0%'),
         ('G', 'Timer overtid (fridager)', f'=SUMIFS(TF_Timer,TF_Overtid,"Ja",{per})', HRS), ('J', 'Herav automatisk LMS', f'={AUTOSUM}', HRS)]
for c, lab, f_, fmt in tiles:
    rap[f'{c}8'] = lab; rap[f'{c}8'].font = F(color='636B74', size=9)
    rap[f'{c}9'] = f_; rap[f'{c}9'].font = F(size=18, bold=True); rap[f'{c}9'].number_format = fmt; rap[f'{c}9'].alignment = Alignment(horizontal='left')

# per aktivitet
rap[f'A{A0 - 1}'] = 'Per aktivitet'; rap[f'A{A0 - 1}'].font = F(size=12, bold=True)
head(rap, A0 - 0, 1, ['Aktivitet', 'Timer', 'Andel'])
A0 += 1; A1 += 1
for i, (an, *_) in enumerate(ACTS):
    r = A0 + i
    rap[f'A{r}'] = an
    rap[f'B{r}'] = f'=SUMIFS(TF_Timer,TF_Akt,A{r},{per})' + (f'+{AUTOSUM}' if an == 'LMS' else '')
    rap[f'C{r}'] = f'=IF($B${A1 + 1}=0,0,B{r}/$B${A1 + 1})'
    rap[f'B{r}'].number_format = HRS; rap[f'C{r}'].number_format = '0%'
    for c in 'ABC': rap[f'{c}{r}'].border = BOX; rap[f'{c}{r}'].font = F()
rap[f'A{A1 + 1}'] = 'Sum'; rap[f'B{A1 + 1}'] = f'=SUM(B{A0}:B{A1})'; rap[f'C{A1 + 1}'] = f'=IF(B{A1 + 1}=0,0,1)'
rap[f'B{A1 + 1}'].number_format = HRS; rap[f'C{A1 + 1}'].number_format = '0%'
for c in 'ABC': rap[f'{c}{A1 + 1}'].font = F(bold=True); rap[f'{c}{A1 + 1}'].fill = HEAD; rap[f'{c}{A1 + 1}'].border = BOX
rap.conditional_formatting.add(f'B{A0}:B{A1}', DataBarRule(start_type='num', start_value=0, end_type='max', color=ACCENT))
# fix tile references (row of LMS and Sum)
rap['A9'] = f'=B{A1 + 1}'; rap['C9'] = f'=B{A0 + 1}'

# per skift (fylles etter per tekniker)
S0 = A1 + 4
rap[f'A{S0 - 1}'] = 'Per skift'; rap[f'A{S0 - 1}'].font = F(size=12, bold=True)
head(rap, S0, 1, ['Skift', 'Timer', 'Herav overtid'])

# LMS per maskin
MC = 5                                            # kolonne E
rap.cell(11, MC, 'LMS per maskin').font = F(size=12, bold=True)
head(rap, 12, MC, ['Maskin'] + WORK + ['Sum'])
m0 = 13
for i in range(len(MACHINES) + 1):
    r = m0 + i; none = i == len(MACHINES)
    rap.cell(r, MC, 'Ingen maskin' if none else f'=INDEX(Maskiner,{i + 1})')
    crit = f'TF_Maskin,""' if none else f'TF_Maskin,{col(MC)}{r}'
    for j in range(len(WORK)):
        rap.cell(r, MC + 1 + j, f'=SUMIFS(TF_Timer,TF_Akt,"LMS",{crit},TF_Arbeid,{col(MC + 1 + j)}$12,{per})')
    rap.cell(r, MC + 1 + len(WORK), f'=SUMIFS(TF_Timer,TF_Akt,"LMS",{crit},{per})' + (f'+{AUTOSUM}' if none else ''))
    for j in range(len(WORK) + 2):
        c = rap.cell(r, MC + j); c.border = BOX; c.font = F(color='636B74' if none and j == 0 else '1B1F24', bold=(j == len(WORK) + 1))
        if j: c.number_format = HRS
mS = m0 + len(MACHINES) + 1
rap.cell(mS, MC, 'Sum')
for j in range(len(WORK) + 1):
    cl = col(MC + 1 + j); rap.cell(mS, MC + 1 + j, f'=SUM({cl}{m0}:{cl}{mS - 1})').number_format = HRS
for j in range(len(WORK) + 2):
    c = rap.cell(mS, MC + j); c.font = F(bold=True); c.fill = HEAD; c.border = BOX
rap.cell(mS + 1, MC, 'Automatisk LMS har ingen maskin og står under «Ingen maskin». Før LMS med maskin og arbeid i Timeføring.').font = F(size=9, color='636B74')

# per tekniker
P0 = mS + 4
rap[f'A{P0 - 1}'] = 'Per tekniker'; rap[f'A{P0 - 1}'].font = F(size=12, bold=True)
rap[f'C{P0 - 1}'] = 'Trykk + / − i margen til venstre for å vise eller skjule teknikerne i et skift.'; rap[f'C{P0 - 1}'].font = F(size=9, color='636B74')
ac = [col(2 + i) for i in range(len(ACTS))]; SUMC = col(2 + len(ACTS)); OTC = col(3 + len(ACTS))
head(rap, P0, 1, ['Tekniker'] + [a[0] for a in ACTS] + ['Sum', 'Herav overtid'], wrap=True)
rap.row_dimensions[P0].height = 30
r = P0 + 1; grp_rows = []
for s, sname, _ in SHIFTS:
    g = r; grp_rows.append(g); t0, t1 = g + 1, g + SLOTS
    rap[f'A{g}'] = f'=INDEX(SkiftNavn,{s})&" · "&COUNTIF(A{t0}:A{t1},"?*")'
    for c in ac + [SUMC, OTC]: rap[f'{c}{g}'] = f'=SUM({c}{t0}:{c}{t1})'
    for c in ['A'] + ac + [SUMC, OTC]:
        rap[f'{c}{g}'].font = F(bold=True); rap[f'{c}{g}'].fill = HEAD; rap[f'{c}{g}'].border = BOX
        if c != 'A': rap[f'{c}{g}'].number_format = HRS
    for n in range(1, SLOTS + 1):
        k = (s - 1) * SLOTS + n; rr = g + n
        rap[f'A{rr}'] = f'=INDEX(Teknikere,{k})&""'
        for i, (an, *_) in enumerate(ACTS):
            f_ = f'SUMIFS(TF_Timer,TF_Tekniker,$A{rr},TF_Akt,{ac[i]}${P0},{per})'
            if an == 'LMS': f_ += f'+SUMIFS(INDEX(GR_Auto,0,{k}),GR_Dato,">="&Fra,GR_Dato,"<="&Til)'
            rap[f'{ac[i]}{rr}'] = f'=IF($A{rr}="","",{f_})'
        rap[f'{SUMC}{rr}'] = f'=IF($A{rr}="","",SUM({ac[0]}{rr}:{ac[-1]}{rr}))'
        rap[f'{OTC}{rr}'] = f'=IF($A{rr}="","",SUMIFS(TF_Timer,TF_Tekniker,$A{rr},TF_Overtid,"Ja",{per}))'
        for c in ['A'] + ac + [SUMC, OTC]:
            rap[f'{c}{rr}'].border = BOX; rap[f'{c}{rr}'].font = F(bold=(c == SUMC))
            if c != 'A': rap[f'{c}{rr}'].number_format = HRS
        rap[f'A{rr}'].alignment = Alignment(indent=1)
    rap.row_dimensions.group(t0, t1, outline_level=1, hidden=True)
    r = t1 + 1
rap[f'A{r}'] = 'Sum'
for c in ac + [SUMC, OTC]: rap[f'{c}{r}'] = '=' + '+'.join(f'{c}{g}' for g in grp_rows)
for c in ['A'] + ac + [SUMC, OTC]:
    rap[f'{c}{r}'].font = F(bold=True); rap[f'{c}{r}'].fill = HEAD; rap[f'{c}{r}'].border = BOX
    if c != 'A': rap[f'{c}{r}'].number_format = HRS
rap.sheet_properties.outlinePr.summaryBelow = False
# per skift fra gruppene
for i, (s, sname, _) in enumerate(SHIFTS):
    rr = S0 + 1 + i
    rap[f'A{rr}'] = f'=INDEX(SkiftNavn,{s})'; rap[f'B{rr}'] = f'={SUMC}{grp_rows[i]}'; rap[f'C{rr}'] = f'={OTC}{grp_rows[i]}'
    for c in 'ABC':
        rap[f'{c}{rr}'].border = BOX; rap[f'{c}{rr}'].font = F()
        if c != 'A': rap[f'{c}{rr}'].number_format = HRS
for c, w in {'A': 22, 'B': 11, 'C': 11, 'D': 11}.items(): rap.column_dimensions[c].width = w
for i in range(5, 16): rap.column_dimensions[col(i)].width = 11
rap.page_setup.orientation = 'landscape'; rap.sheet_properties.pageSetUpPr.fitToPage = True
rap.page_setup.fitToWidth = 1; rap.page_setup.fitToHeight = 0

for ws in (kal, tf, rap, opp):
    ws.sheet_view.showGridLines = ws is not kal and ws is not rap
    for row in ws.iter_rows():
        for c in row:
            if c.font is None or c.font.name != FONT: c.font = F()
kal.sheet_properties.tabColor = ACCENT
wb.calculation.fullCalcOnLoad = True
wb.save(OUT)
print('Skrev', OUT)
